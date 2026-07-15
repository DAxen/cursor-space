# -*- coding: utf-8 -*-
"""
生成两页PPT:
  第一页:知识工程试点落地情况 —— Lite团队OH中枢特性试点
  第二页:质量工程建设整体思路 —— 基于MTG建模的LLT用例自动生成

风格要求:
- 背景: 纯白
- 字体: 微软雅黑
- 大标题: #C00000, 24号, 加粗, 左对齐
- 副标题: 字体 #000000, 14号, 背景框 #F2F2F2
- 正文: #000000, 10~12号
- 正文强调: #0000FF
排版:
- 最上方大标题(左对齐) -> 副标题 -> 总结/结论信息 -> 正文/图示
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ---------- 常量 ----------
FONT_NAME = "微软雅黑"
COLOR_TITLE = RGBColor(0xC0, 0x00, 0x00)
COLOR_TEXT = RGBColor(0x00, 0x00, 0x00)
COLOR_EMPH = RGBColor(0x00, 0x00, 0xFF)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_SUBBG = RGBColor(0xF2, 0xF2, 0xF2)
COLOR_ACCENT = RGBColor(0xC0, 0x00, 0x00)
COLOR_CARD_BORDER = RGBColor(0xD9, 0xD9, 0xD9)
COLOR_FLOW = RGBColor(0xC0, 0x00, 0x00)

# 16:9 幻灯片
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN_L = Inches(0.5)
CONTENT_W = SLIDE_W - Inches(1.0)


def set_run_font(run, size, color, bold=False, italic=False):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set('typeface', FONT_NAME)


def set_white_background(slide, prs):
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_WHITE
    bg.line.fill.background()
    bg.shadow.inherit = False
    sp = bg._element
    sp.getparent().remove(sp)
    slide.shapes._spTree.insert(2, sp)
    return bg


def add_title(slide, text, top=Inches(0.32)):
    box = slide.shapes.add_textbox(MARGIN_L, top, CONTENT_W, Inches(0.6))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = text
    set_run_font(r, 24, COLOR_TITLE, bold=True)
    return box


def add_title_rule(slide, top):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, MARGIN_L, top, CONTENT_W, Pt(2.2)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_ACCENT
    line.line.fill.background()
    line.shadow.inherit = False
    return line


def add_subtitle(slide, text, top, height=Inches(0.4)):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN_L, top,
                                 CONTENT_W, height)
    box.fill.solid()
    box.fill.fore_color.rgb = COLOR_SUBBG
    box.line.fill.background()
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.12)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = text
    set_run_font(r, 14, COLOR_TEXT, bold=True)
    return box


def add_conclusion(slide, segments, top, height=Inches(0.55), size=11.5):
    box = slide.shapes.add_textbox(MARGIN_L, top, CONTENT_W, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_top = 0
    add_para(tf, segments, size=size, new=False, space_after=0)
    return box


def add_card(slide, left, top, width, height, title, title_size=11.5):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    card.adjustments[0] = 0.05
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_WHITE
    card.line.color.rgb = COLOR_CARD_BORDER
    card.line.width = Pt(1.0)
    card.shadow.inherit = False
    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.14)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.09)
    tf.margin_bottom = Inches(0.08)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = title
    set_run_font(r, title_size, COLOR_TITLE, bold=True)
    p.space_after = Pt(3)
    return card, tf


def add_para(tf, segments, size=11, bullet=False, level=0,
             space_before=1.5, space_after=1.5, new=True, line_spacing=None):
    p = tf.add_paragraph() if new else tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.level = level
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if line_spacing:
        p.line_spacing = line_spacing
    if bullet:
        r = p.add_run()
        r.text = "• " if level == 0 else "– "
        set_run_font(r, size, COLOR_TEXT)
    for text, emph in segments:
        r = p.add_run()
        r.text = text
        set_run_font(r, size, COLOR_EMPH if emph else COLOR_TEXT, bold=emph)
    return p


def add_metric_card(slide, left, top, width, height, number, label,
                    num_size=18, label_size=9.5):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    card.adjustments[0] = 0.1
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_SUBBG
    card.line.color.rgb = COLOR_CARD_BORDER
    card.line.width = Pt(1.0)
    card.shadow.inherit = False
    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = number
    set_run_font(r, num_size, COLOR_EMPH, bold=True)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(1)
    r2 = p2.add_run()
    r2.text = label
    set_run_font(r2, label_size, COLOR_TEXT)
    return card


def add_flow_step(slide, left, top, width, height, index, title, details):
    """流程步骤框:序号 + 标题 + 说明。"""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    card.adjustments[0] = 0.08
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_WHITE
    card.line.color.rgb = COLOR_FLOW
    card.line.width = Pt(1.4)
    card.shadow.inherit = False
    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.1)
    tf.margin_bottom = Inches(0.08)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = f"STEP {index}  "
    set_run_font(r, 10, COLOR_FLOW, bold=True)
    r2 = p.add_run()
    r2.text = title
    set_run_font(r2, 11, COLOR_TEXT, bold=True)
    p.space_after = Pt(2)
    for seg in details:
        add_para(tf, seg, size=9.5, space_before=1, space_after=1,
                 line_spacing=1.0)
    return card


def add_arrow(slide, left, top, width, height):
    ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    ar.adjustments[0] = 0.55
    ar.adjustments[1] = 0.55
    ar.fill.solid()
    ar.fill.fore_color.rgb = COLOR_FLOW
    ar.line.fill.background()
    ar.shadow.inherit = False
    return ar


def add_footer(slide, text, prs):
    box = slide.shapes.add_textbox(MARGIN_L, prs.slide_height - Inches(0.4),
                                   CONTENT_W, Inches(0.3))
    tf = box.text_frame
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = text
    set_run_font(r, 9, RGBColor(0x80, 0x80, 0x80))


# ============================================================
def build_page1(prs, blank):
    s = prs.slides.add_slide(blank)
    set_white_background(s, prs)
    add_title(s, "知识工程试点落地情况 —— Lite团队OH中枢特性试点")
    add_title_rule(s, Inches(0.95))
    add_subtitle(s, "试点方案、效果评价、关键数据与下一步计划", Inches(1.06))

    add_conclusion(
        s,
        [("结论:", True),
         ("基于 okl 工具(参考 LLM Wiki 方法论),逆向分析存量代码/文档并正向投喂设计文档,"
          "在 OH 中枢代码仓构建知识库;对代码中", False),
         ("无法映射的概念", True),
         (",带知识库问答效果明显优于无知识库,回答更贴合设计文档;对", False),
         ("可推导知识", True),
         ("差距不大。", False)],
        top=Inches(1.54), height=Inches(0.62), size=11)

    # 关键指标卡
    metric_top = Inches(2.2)
    metric_h = Inches(0.82)
    mgap = Inches(0.22)
    mw = (CONTENT_W - 2 * mgap) / 3
    metrics = [
        ("150 kloc", "逆向分析代码规模"),
        ("4 + 69 篇", "投喂:需求分析/功能设计 + 实现设计(MD)"),
        ("215 篇", "知识库摄取 Wiki 文档总数"),
    ]
    for i, (num, label) in enumerate(metrics):
        add_metric_card(s, MARGIN_L + i * (mw + mgap), metric_top, mw,
                        metric_h, num, label)

    # 2x2 内容卡片
    grid_top = Inches(3.2)
    row_gap = Inches(0.18)
    col_gap = Inches(0.28)
    card_w = (CONTENT_W - col_gap) / 2
    card_h = Inches(1.72)
    r2_top = grid_top + card_h + row_gap
    lx = MARGIN_L
    rx = MARGIN_L + card_w + col_gap

    _, a = add_card(s, lx, grid_top, card_w, card_h, "一、内容描述(做了什么)")
    add_para(a, [("以 ", False), ("okl 工具", True),
                 ("(参考 LLM Wiki 方法论)为核心开展试点;", False)],
             bullet=True, size=10)
    add_para(a, [("逆向分析", True), (" 存量代码与文档,提炼隐性知识;", False)],
             bullet=True, size=10)
    add_para(a, [("正向投喂", True), (" 设计文档,补齐代码无法表达的概念;", False)],
             bullet=True, size=10)
    add_para(a, [("在 ", False), ("OH 中枢代码仓", True),
                 (" 内构建可检索、可问答的知识库。", False)],
             bullet=True, size=10)

    _, b = add_card(s, rx, grid_top, card_w, card_h, "二、效果评价(主观)")
    add_para(b, [("代码中", False), ("无法映射的概念", True),
                 (":带库问答", False), ("明显优于", True), ("无库;", False)],
             bullet=True, size=10)
    add_para(b, [("回答", False), ("贴合设计文档中的概念", True),
                 (",解释更准确完整;", False)], bullet=True, size=10)
    add_para(b, [("代码中", False), ("可直接推导的知识", True),
                 (":带/不带库", False), ("差距不大", True), (";", False)],
             bullet=True, size=10)
    add_para(b, [("判断:知识库在", False), ("补齐设计意图类概念", True),
                 ("上价值最突出。", False)], bullet=True, size=10)

    _, c = add_card(s, lx, r2_top, card_w, card_h, "三、关键数据 · 摄取产出明细(共 215 篇)")
    add_para(c, [("流程页 ", False), ("154 篇", True),
                 (":含 ", False), ("24 条流程", True), (" 的深挖子页;", False),
                 ("  子模块页 ", False), ("34 篇", True), (";", False)],
             bullet=True, size=10)
    add_para(c, [("全局页 ", False), ("10 篇", True),
                 (":业务域、契约、用例、架构;", False)], bullet=True, size=10)
    add_para(c, [("仓级页 ", False), ("17 篇", True),
                 (":overview、架构、数据模型等补充;", False)], bullet=True, size=10)
    add_para(c, [("投喂来源:", False), ("4 篇", True),
                 (" wxalm 需求分析与功能设计 + ", False), ("69 篇", True),
                 (" 实现设计(markdown)。", False)], bullet=True, size=10)

    _, d = add_card(s, rx, r2_top, card_w, card_h, "四、下一步计划")
    add_para(d, [("目录规整", True), (":梳理知识库结构,提升可读性与检索效率;", False)],
             bullet=True, size=10)
    add_para(d, [("引入", False), ("知识库测评问题集", True),
                 (",量化评估问答质量;", False)], bullet=True, size=10)
    add_para(d, [("在", False), ("新需求", True), ("中应用知识库", False),
                 ("辅助编码", True), (",验证工程价值。", False)],
             bullet=True, size=10)

    add_footer(s, "第 1 页 / 共 2 页", prs)


def build_page2(prs, blank):
    s = prs.slides.add_slide(blank)
    set_white_background(s, prs)
    add_title(s, "质量工程建设整体思路 —— 基于MTG建模的LLT用例自动生成")
    add_title_rule(s, Inches(0.95))
    add_subtitle(s, "内容描述、TDD关键流程、关键点与进展计划", Inches(1.06))

    add_conclusion(
        s,
        [("思路:", True),
         ("结合 ", False), ("MTG(Model Testing Generator)", True),
         (" 实践,将绘制好的", False), ("活动图", True),
         ("结合 AI 自动生成用例数据,自动生成测试用例,", False),
         ("使能 TDD 开发流程", True), ("。", False)],
        top=Inches(1.54), height=Inches(0.5), size=11)

    # ---- TDD 关键流程图(4 步) ----
    flow_top = Inches(2.12)
    flow_h = Inches(1.62)
    arrow_w = Inches(0.34)
    n = 4
    step_w = (CONTENT_W - (n - 1) * arrow_w) / n
    steps = [
        ("绘制活动图(关键)",
         [[("基于 ", False), ("AR 描述的业务场景", True), ("绘制;", False)],
          [("入口为组件/对外模块", False), ("API 入口", True), ("(plantuml)", False)]]),
        ("MTG 生成用例设计",
         [[("基于活动图 + 已有 ", False), ("MTG 用例生成工具", True), (";", False)],
          [("输出用例设计", False), ("(markdown)", True)]]),
        ("生成用例实现代码",
         [[("基于用例设计输出", False), ("(markdown)", True)],
          [("自动生成", False), ("用例实现代码", True), ("。", False)]]),
        ("开发并跑通用例",
         [[("开发功能代码,", False)],
          [("将用例", False), ("执行通过", True), (",完成 TDD 闭环。", False)]]),
    ]
    x = MARGIN_L
    for i, (title, details) in enumerate(steps, 1):
        add_flow_step(s, x, flow_top, step_w, flow_h, i, title, details)
        x += step_w
        if i < n:
            add_arrow(s, x, flow_top + flow_h / 2 - Inches(0.16),
                      arrow_w, Inches(0.32))
            x += arrow_w

    # ---- 下方三块:关键点 / 当前进展 / 下一步计划 ----
    low_top = Inches(3.95)
    low_h = Inches(2.9)
    col_gap = Inches(0.26)
    left_w = (CONTENT_W - col_gap) * 0.52
    right_w = (CONTENT_W - col_gap) * 0.48

    _, k = add_card(s, MARGIN_L, low_top, left_w, low_h, "关键点")
    add_para(k, [("活动图", True), ("尽量", False), ("不涉及代码实现细节", True),
                 (",不映射代码元素,基于", False), ("业务场景", True), ("描述;", False)],
             bullet=True, size=10.5)
    add_para(k, [("因 ", False), ("API 接口稳定", True),
                 (",据此生成的用例代码也", False), ("稳定", True),
                 (",稳定性", False), ("高于传统 UT", True),
                 (",不随代码实现频繁变化;", False)], bullet=True, size=10.5)
    add_para(k, [("活动图与用例设计输出是", False), ("结构化", True),
                 ("的,利于 AI ", False), ("正向生成用例", True),
                 (",而非基于代码实现反推。", False)], bullet=True, size=10.5)

    rx = MARGIN_L + left_w + col_gap
    _, p = add_card(s, rx, low_top, right_w, low_h, "当前进展 & 下一步计划")
    add_para(p, [("【当前进展】", True)], size=10.5, space_before=1)
    add_para(p, [("已讨论明确", False), ("整体方案思路", True), (";", False)],
             bullet=True, size=10.5)
    add_para(p, [("已选定", False), ("AI 团队 agent 框架需求", True),
                 (" 与 ", False), ("Lite 团队 OH 中枢需求", True),
                 (" 作为试点项目。", False)], bullet=True, size=10.5)
    add_para(p, [("【下一步计划】", True)], size=10.5, space_before=6)
    add_para(p, [("在试点项目中选取", False), ("存量需求", True),
                 (",绘制 ", False), ("MTG 活动图", True),
                 (",并尝试", False), ("测试用例生成", True), ("。", False)],
             bullet=True, size=10.5)

    add_footer(s, "第 2 页 / 共 2 页", prs)


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    build_page1(prs, blank)
    build_page2(prs, blank)
    out = "知识工程与质量工程建设_汇报PPT.pptx"
    prs.save(out)
    print("saved:", out)


if __name__ == "__main__":
    build()
